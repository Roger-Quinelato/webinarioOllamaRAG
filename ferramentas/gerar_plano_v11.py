import copy
import re
import sys
from pathlib import Path

import docx

RAIZ = Path(__file__).resolve().parents[1]
ORIGEM = Path(sys.argv[1]) if len(sys.argv) > 1 else next(
    (Path.home() / "Downloads").glob("Plano_Aula_2-*v1.0.docx"))
DESTINO = RAIZ / "docs" / "Plano_Aula_2-Construindo um Assistente com RAG, Ollama e Streamlit_v1.1.docx"

documento = docx.Document(str(ORIGEM))
paragrafos = documento.paragraphs


def trocar(paragrafo, texto):
    primeiro = copy.deepcopy(paragrafo.runs[0]._element) if paragrafo.runs else None
    for filho in list(paragrafo._element):
        if filho.tag != docx.oxml.ns.qn("w:pPr"):
            paragrafo._element.remove(filho)
    if primeiro is None:
        paragrafo.add_run(texto)
        return
    paragrafo._element.append(primeiro)
    paragrafo.runs[-1].text = texto


def substituir_lista(primeiro, ultimo, itens):
    modelo = paragrafos[primeiro]
    for indice in range(primeiro + 1, ultimo + 1):
        elemento = paragrafos[indice]._element
        elemento.getparent().remove(elemento)
    trocar(modelo, itens[0])
    anterior = modelo
    for item in itens[1:]:
        novo = copy.deepcopy(modelo._element)
        anterior._element.addnext(novo)
        anterior = docx.text.paragraph.Paragraph(novo, modelo._parent)
        trocar(anterior, item)


def definir_celula(celula, texto):
    trocar(celula.paragraphs[0], texto)
    for extra in celula.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


trocar(paragrafos[4], "Versão 1.1 – Atual  |  Setembro de 2026  |  Turmas: 21/09/2026 (CIIA) e 28/09/2026 (aberta), 16h, online")

cabecalho_tabela = documento.tables[0]
definir_celula(cabecalho_tabela.rows[2].cells[1],
               "Ollama: https://ollama.com  ·  ChromaDB: https://docs.trychroma.com  ·  "
               "Streamlit: https://docs.streamlit.io  ·  SHAP: https://shap.readthedocs.io  ·  "
               "VS Code: https://code.visualstudio.com")

substituir_lista(44, 46, [
    "1. LEWIS, P. et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS, 2020.",
    "2. KARPUKHIN, V. et al. Dense Passage Retrieval for Open-Domain Question Answering. EMNLP, 2020.",
    "3. GAO, Y. et al. Retrieval-Augmented Generation for Large Language Models: A Survey. arXiv:2312.10997, 2023.",
    "4. ES, S. et al. RAGAS: Automated Evaluation of Retrieval Augmented Generation. arXiv:2309.15217, 2023.",
    "5. ASAI, A. et al. Self-RAG: Learning to Retrieve, Generate, and Critique through Self-Reflection. ICLR, 2024.",
    "6. LIU, N. F. et al. Lost in the Middle: How Language Models Use Long Contexts. TACL, 2024.",
])

substituir_lista(39, 41, [
    "Indexação dos artigos: extração de texto por página com pypdf, chunking com sobreposição, embeddings com "
    "bge-m3 via Ollama e persistência no ChromaDB com metadados escritos à mão (ano, tema, idioma) e resumos "
    "gerados pelo LLM; demonstração da extração de metadados pelo próprio LLM comparada ao CSV.",
    "Retrieval top-k com inspeção das distâncias, filtros de metadados, busca cross-lingual e busca em dois "
    "estágios (resumos → trechos); explicação do retrieval com SHAP e contribuição de cada trecho à resposta "
    "(valores de Shapley pré-computados).",
    "Montagem do prompt enriquecido e comparação das mesmas perguntas com e sem contexto, evidenciando o efeito "
    "do RAG sobre alucinações; integração com qwen2.5 servido pelo Ollama em streaming e chatbot em Streamlit; "
    "encerramento com avaliação no estilo RAGAS, reranking e escalabilidade.",
])

substituir_lista(30, 32, [
    "Hardware/Software: Python 3.10+, VS Code (extensões Python e Jupyter), Ollama local; 8 GB de RAM (LLM em CPU).",
    "Modelos e bibliotecas: bge-m3 (embeddings), qwen2.5:3b e qwen2.5:1.5b (chat), ollama, chromadb, pypdf, shap, "
    "streamlit — sem LangChain/LlamaIndex, para expor o mecanismo do RAG.",
    "Material didático: repositório com notebook da aula (saídas salvas), scripts numerados por bloco, app "
    "Streamlit, README de instalação, vídeo de instalação no YouTube do CIIA e 6 artigos do arXiv como base documental.",
])

substituir_lista(24, 26, [
    "Aula expositiva dialogada com demonstração ao vivo do pipeline completo pelo facilitador.",
    "Suporte online no chat por um segundo facilitador, com respostas prontas para os erros mais comuns.",
    "Replicação posterior pelos participantes a partir do repositório, do notebook com saídas salvas e do vídeo "
    "de instalação (os downloads de modelos não são feitos durante a aula).",
])

substituir_lista(16, 20, [
    "Recapitulação do pipeline RAG e do ambiente local (Ollama, modelos, .venv);",
    "Indexação: chunking, embeddings, ChromaDB e metadados (manuais e extraídos por LLM);",
    "Retrieval semântico: seleção top-k, filtros de metadados e busca cross-lingual;",
    "Busca em dois estágios (resumos → trechos);",
    "Explicabilidade do retrieval com SHAP;",
    "Prompt augmentation e comparação de respostas com e sem contexto;",
    "Conexão com LLM local (Ollama) em streaming, com citação de fontes;",
    "Chatbot interativo com Streamlit;",
    "Avaliação (RAGAS), reranking e escalabilidade.",
])

substituir_lista(9, 13, [
    "Indexar artigos científicos em um banco vetorial local com metadados úteis para filtragem;",
    "Implementar e analisar o retrieval top-k, com filtros de metadados e busca em dois estágios;",
    "Explicar o comportamento do retrieval com SHAP;",
    "Montar prompts enriquecidos com contexto recuperado (prompt augmentation);",
    "Integrar a recuperação com um LLM local servido pelo Ollama;",
    "Avaliar comparativamente respostas com e sem contexto recuperado;",
    "Desenvolver uma interface conversacional interativa com Streamlit.",
])

cronograma = [
    ("1", "Recapitulando o pipeline RAG e o ambiente local", "10 min"),
    ("2", "Indexação: PDFs → chunks → embeddings → ChromaDB, metadados manuais × LLM e resumos", "18 min"),
    ("3", "Retrieval top-k na prática: k, distâncias, filtros e cross-lingual", "15 min"),
    ("3b", "Busca em dois estágios: resumos → trechos", "7 min"),
    ("4", "SHAP: explicando o retrieval (ao vivo) e Shapley dos trechos (pré-computado)", "12 min"),
    ("5", "Prompt augmentation: respostas com e sem contexto", "12 min"),
    ("6", "Integração com LLM local usando Ollama (streaming e fontes)", "12 min"),
    ("7", "Construção do chatbot com Streamlit", "15 min"),
    ("8", "Avaliação (RAGAS), reranking e próximos passos", "8 min"),
    ("—", "Folga para imprevistos", "5 min"),
    ("", "Total:", "1h54"),
]
tabela = documento.tables[1]
linha_modelo = tabela.rows[1]._tr
for linha in list(tabela.rows)[1:]:
    tabela._tbl.remove(linha._tr)
for bloco, atividade, duracao in cronograma:
    nova = copy.deepcopy(linha_modelo)
    tabela._tbl.append(nova)
    celulas = tabela.rows[-1].cells
    for celula, texto in zip(celulas, (bloco, atividade, duracao)):
        definir_celula(celula, texto)

anexos = [p for p in documento.paragraphs if re.match(r"Anexo \d", p.text.strip())]
textos_anexos = [
    "Anexo 1: Checklist de preparação do ambiente — (1) instalar o Ollama; (2) opcional: definir OLLAMA_MODELS; "
    "(3) ollama pull bge-m3, qwen2.5:3b e qwen2.5:1.5b; (4) python -m venv .venv e pip install -r requirements.txt; "
    "(5) registrar o kernel webinario-rag; (6) python scripts/00_checar_ambiente.py até “Ambiente pronto.”; "
    "(7) scripts/01_preparar_corpus.py e scripts/02_indexar.py. Detalhes no README.md do repositório.",
    "Anexo 2: Estrutura do material — config.py (modelos e parâmetros), rag.py (funções do pipeline), "
    "metadados.csv, scripts/00–07 (um por bloco), webinario_rag.ipynb (notebook da aula com saídas salvas), "
    "app.py (Streamlit), opcional/ (Shapley dos trechos e avaliação no estilo RAGAS), resultados/ (saídas "
    "pré-computadas) e docs/ (roteiro, troubleshooting, medições e verificação).",
    "Anexo 3: Roteiro do facilitador e planos B — fala, demonstração, checkpoint e plano B por bloco em "
    "docs/roteiro_facilitador.md; tempos medidos na máquina de demonstração em docs/medicoes.md; erros comuns "
    "e respostas para o chat em docs/troubleshooting.md.",
]
for paragrafo, texto in zip(anexos, textos_anexos):
    trocar(paragrafo, texto)

documento.save(str(DESTINO))
print(f"Salvo em {DESTINO}")
